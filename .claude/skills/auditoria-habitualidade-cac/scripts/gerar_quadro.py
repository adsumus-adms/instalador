#!/usr/bin/env python3
"""Gera o quadro de habitualidade (.xlsx) a partir de um JSON.

Uso: python3 -I scripts/gerar_quadro.py entrada.json saida.xlsx
Esquema do JSON: ver references/quadro-schema.md
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL = "1F3F73"
FILL_HEAD = PatternFill("solid", fgColor=AZUL)
FILL = {
    "sim": PatternFill("solid", fgColor="D9EFE3"),
    "nao": PatternFill("solid", fgColor="F8D7D7"),
    "confirmar": PatternFill("solid", fgColor="FBF1D9"),
}
LADO = Side(style="thin", color="C9CED8")
BORDA = Border(left=LADO, right=LADO, top=LADO, bottom=LADO)
FONTE = "Arial"


def cabecalho(ws, linha, titulos):
    for i, t in enumerate(titulos, 1):
        c = ws.cell(row=linha, column=i, value=t)
        c.font = Font(name=FONTE, bold=True, color="FFFFFF", size=10)
        c.fill = FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDA


def linha(ws, r, valores, status_col=None):
    for i, v in enumerate(valores, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(name=FONTE, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        c.border = BORDA
    if status_col:
        v = str(valores[status_col - 1] or "").lower()
        chave = "sim" if v.startswith("sim") or v == "ok" else "nao" if v.startswith("n") else "confirmar"
        ws.cell(row=r, column=status_col).fill = FILL[chave]
        ws.cell(row=r, column=status_col).font = Font(name=FONTE, size=10, bold=True)


def larguras(ws, ws_larguras):
    for i, w in enumerate(ws_larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def titulo(ws, texto, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=texto)
    c.font = Font(name=FONTE, bold=True, size=14, color=AZUL)
    ws.row_dimensions[1].height = 24


def main():
    entrada, saida = sys.argv[1], sys.argv[2]
    d = json.load(open(entrada, encoding="utf-8"))
    cac = d["cac"]
    wb = Workbook()

    # Aba 1: Resumo
    ws = wb.active
    ws.title = "Resumo"
    ws.sheet_view.showGridLines = False
    titulo(ws, "Habitualidade: " + cac["nome"], 7)
    r = 3
    for rot, chave in [("CR", "cr"), ("Emissão do CR", "emissao_cr"), ("Padrão", "padrao"),
                       ("Filiação", "filiacao"), ("Data da declaração analisada", "declaracao_data"),
                       ("Situação da análise", "status_analise")]:
        if cac.get(chave):
            ws.cell(row=r, column=1, value=rot).font = Font(name=FONTE, bold=True, size=10)
            ws.cell(row=r, column=2, value=cac[chave]).font = Font(name=FONTE, size=10)
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
            r += 1
    r += 1
    ws.cell(row=r, column=1, value="Resumo por ciclo").font = Font(name=FONTE, bold=True, size=11, color=AZUL)
    r += 1
    cabecalho(ws, r, ["Ciclo", "Início", "Fim", "Treinos confirmados", "Meta", "Cumpre?", "Observação"])
    ws.row_dimensions[r].height = 30
    for c in d["ciclos"]:
        r += 1
        linha(ws, r, [c["ciclo"], c["inicio"], c["fim"], c["confirmados"], c.get("meta", 8),
                      c["cumpre"], c.get("obs", "")], status_col=6)
    r += 2
    ws.cell(row=r, column=1, value="Armas").font = Font(name=FONTE, bold=True, size=11, color=AZUL)
    r += 1
    cabecalho(ws, r, ["Sigma", "Calibre", "Classificação na declaração", "Grupo (1 a 6)", "Eventos", "", "Observação"])
    ws.row_dimensions[r].height = 30
    for a in d.get("armas", []):
        r += 1
        linha(ws, r, [a.get("sigma", ""), a.get("calibre", ""), a.get("classificacao", ""),
                      a.get("grupo", ""), a.get("eventos", ""), "", a.get("obs", "")])
    larguras(ws, [30, 14, 24, 20, 14, 14, 60])

    # Aba 2: Detalhe dia a dia
    ws2 = wb.create_sheet("Dia a dia")
    ws2.sheet_view.showGridLines = False
    titulo(ws2, "Sistema × Folha × Declaração", 9)
    cabecalho(ws2, 3, ["Ciclo", "Data", "Sigma", "Calibre", "Qtd munição", "Evento",
                       "Sistema", "Folha", "Declaração", "Situação", "Fonte"][:11])
    r = 3
    for x in d.get("detalhe", []):
        r += 1
        linha(ws2, r, [x.get("ciclo", ""), x["data"], x.get("sigma", ""), x.get("calibre", ""),
                       x.get("qtd", ""), x.get("evento", ""), x.get("sistema", ""), x.get("folha", ""),
                       x.get("declaracao", ""), x.get("situacao", ""), x.get("fonte", "")])
    larguras(ws2, [10, 18, 12, 11, 12, 18, 14, 14, 14, 30, 26])
    ws2.freeze_panes = "A4"

    # Aba 3: Pendências
    ws3 = wb.create_sheet("Pendências")
    ws3.sheet_view.showGridLines = False
    titulo(ws3, "Pendências e divergências", 2)
    cabecalho(ws3, 3, ["Nº", "Descrição"])
    r = 3
    for i, p in enumerate(d.get("pendencias", []), 1):
        r += 1
        linha(ws3, r, [i, p])
    larguras(ws3, [6, 110])

    for w in (ws, ws2, ws3):
        w.page_setup.orientation = "landscape"
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 0
        w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(saida)
    print("OK", saida)


if __name__ == "__main__":
    main()
